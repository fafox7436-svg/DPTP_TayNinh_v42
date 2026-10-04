"""Portable, explicit downloads; no automatic server-side persistence of user data."""
from __future__ import annotations

import io
import json

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from data_io import DAYS, ENERGY, HUMIDITY, MONTH, OUTAGE, TEMP, YEAR, InputError


def excel_bytes(sheets: dict[str, pd.DataFrame]) -> bytes:
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for title, frame in sheets.items():
            safe = frame.copy()
            for col in safe.select_dtypes(include=["object", "string"]).columns:
                safe[col] = safe[col].map(lambda v: "'" + v if isinstance(v, str) and v.lstrip().startswith(("=", "+", "-", "@")) else v)
            safe.to_excel(writer, sheet_name=title[:31], index=False)
            ws = writer.sheets[title[:31]]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="123D64")
                cell.font = Font(name="Calibri", color="FFFFFF", bold=True)
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.row_dimensions[1].height = 34
            for cells in ws.columns:
                values = list(cells)
                ws.column_dimensions[values[0].column_letter].width = min(48, max(16, len(str(values[0].value)) + 3))
                for cell in values[1:]:
                    cell.font = Font(name="Calibri", size=11)
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = "#,##0.00;[Red]-#,##0.00"
    return buffer.getvalue()


def input_template() -> bytes:
    return excel_bytes({
        "Lich_su": pd.DataFrame(columns=[YEAR, MONTH, ENERGY, TEMP, HUMIDITY, DAYS, OUTAGE]),
        "Du_bao": pd.DataFrame(columns=[YEAR, MONTH, TEMP, HUMIDITY, DAYS, OUTAGE]),
        "Lich_nghi": pd.DataFrame(columns=["Ngày", "Loại"]),
        "Huong_dan": pd.DataFrame({"Nội dung": [
            "Điền các sheet, sau đó có thể tải cùng tệp vào cả hai ô Lịch sử và Dự báo; chọn đúng sheet.",
            "Một dòng/tháng; lịch sử tối thiểu 24 tháng liên tục. Năm, Tháng là số nguyên.",
            "Tổng thương phẩm: toàn bộ hai tệp dùng cùng đơn vị kWh hoặc triệu kWh đã chọn trong ứng dụng.",
            "Nếu không dùng nhiệt độ/độ ẩm, xóa hẳn cột đó ở cả hai sheet; không để cột rỗng.",
            "Số ngày: có thể xóa cột để tính từ lịch. Số ngày cắt điện: có thể xóa cột để mặc định 0.",
            "Nếu có cột Số ngày cắt điện, điền số ngày tương đương mất điện trên toàn phạm vi, không dùng số vụ.",
            "Lịch_nghi: ngày cụ thể dạng ô Date hoặc YYYY-MM-DD; loại Lễ hoặc Tết, gồm cả ngày nghỉ bù áp dụng.",
            "Không có lịch nghỉ mặc định chính thức. Kiểm tra lịch và xác nhận giả định trước khi dự báo.",
        ]}),
    })


def scenario_json(signature: str, model: str, scenarios: pd.DataFrame) -> bytes:
    return json.dumps({"schema_version": 1, "run_signature": signature, "model": model,
                       "scenarios": scenarios.to_dict(orient="records")},
                      ensure_ascii=False, allow_nan=False, indent=2).encode("utf-8")


def restore_scenarios(data: bytes, signature: str, expected: pd.DataFrame,
                      models: list[str]) -> tuple[str, pd.DataFrame]:
    if len(data) > 1_000_000:
        raise InputError("Tệp phương án vượt 1 MB.")
    try:
        obj = json.loads(data.decode("utf-8"))
        if obj["schema_version"] != 1 or obj["run_signature"] != signature:
            raise InputError("Phương án thuộc dữ liệu hoặc cấu hình khác. Hãy dùng đúng hai tệp và cấu hình gốc.")
        if obj["model"] not in models:
            raise InputError("Mô hình của phương án không có trong lần dự báo này.")
        frame = pd.DataFrame(obj["scenarios"])
        if set(frame.columns) != set(expected.columns) or frame["Tháng"].tolist() != expected["Tháng"].tolist():
            raise InputError("Cấu trúc hoặc danh sách tháng của phương án không hợp lệ.")
        return obj["model"], frame[expected.columns]
    except InputError:
        raise
    except (KeyError, TypeError, ValueError, UnicodeDecodeError) as exc:
        raise InputError("Tệp JSON phương án không hợp lệ.") from exc
